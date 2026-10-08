import json
from pathlib import Path

import pytest

from scripts import test_omission_audit as audit

PUBLIC = Path(audit.PUBLIC)


def published():
    return (json.loads((PUBLIC / "results.json").read_text()),
            json.loads((PUBLIC / "frozen-manifest-public.json").read_text()))


def test_frame_matches_committed_file_and_expected_groups():
    results, manifest = published()
    built = audit.frame(results, manifest)
    committed = json.loads((audit.AUDIT / "audit-frame.json").read_text())
    assert built == committed
    assert built["counts"] == audit.EXPECTED
    assert len(built["items"]) == 104
    assert len({i["call_id"] for i in built["items"]}) == 104


def test_frame_groups_follow_public_verdicts():
    results, manifest = published()
    for item in audit.frame(results, manifest)["items"]:
        if item["group"] == "R":
            assert item["reference_verdict"] in audit.ALARMS
        else:
            assert item["reference_verdict"] == "quiet" and item["shown_mutant_verdict"] == "quiet"
            assert item["source"] != "spec_complete"


def test_frame_rejects_drift():
    results, manifest = published()
    results = {**results, "rows": [r for r in results["rows"] if r["source"] != "spec_complete"]}
    with pytest.raises(ValueError, match="drift"):
        audit.frame(results, manifest)


def test_blind_order_mixes_sources_and_is_deterministic():
    items = audit.AUDIT.joinpath("audit-frame.json")
    items = json.loads(items.read_text())["items"]
    first, second = audit.blind_order(items), audit.blind_order(items)
    assert first == second
    assert [i["item_id"] for i in first] == [f"T{n:03d}" for n in range(1, 105)]
    assert len({i["source"] for i in first[:15]}) > 1


def test_mark_rule_wraps_the_removed_span():
    assert audit.mark_rule("Save. Only admins may delete. Done.", "Save. Done.") == "Save. [[Only admins may delete. ]]Done."
    with pytest.raises(ValueError):
        audit.mark_rule("abc", "xyz")


def fake_study(tmp_path: Path, frame: dict) -> Path:
    study = tmp_path / "study"
    results, _ = published()
    (study).mkdir()
    (study / "results.json").write_text(json.dumps(results))
    for item in frame["items"]:
        call = study / "calls" / item["call_id"]
        call.mkdir(parents=True, exist_ok=True)
        (call / "suite.cjs").write_text(f"// suite {item['call_id']}\nassert.equal(state.rule, true);\n")
        for slot, outcome in ((item["reference_slot"], "assertion_failure"), (item["shown_mutant_slot"], "pass")):
            d = study / "execution" / item["call_id"] / slot
            d.mkdir(parents=True, exist_ok=True)
            (d / "report.json").write_text(json.dumps({"status": "complete", "tests": [
                {"name": "rule", "outcome": outcome, "message": "expected true"}]}))
    return study


def fill(path: Path, out: Path, label_for, quote="assert.equal(state.rule, true);"):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    for group, title in audit.SHEETS.items():
        ws = wb[title]
        header = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2):
            item = row[0].value
            row[header.index("categoria")].value = label_for(group, item)
            row[header.index("menciona_regra")].value = "sim"
            row[header.index("citacao")].value = quote
            row[header.index("palpite_fonte")].value = "nao_sei"
    wb.save(out)


def test_sheet_and_score_end_to_end(tmp_path):
    pytest.importorskip("openpyxl")
    frame = json.loads((audit.AUDIT / "audit-frame.json").read_text())
    study = fake_study(tmp_path, frame)
    out = tmp_path / "audit"
    built = audit.build_sheet(study, out, frame)
    assert built["items"] == 104 and built["groups"] == {"P": 52, "R": 52}
    with pytest.raises(FileExistsError):
        audit.build_sheet(study, out, frame)
    sheet_text = (out / "auditoria-testes-codificador_a.xlsx").read_bytes()
    assert b"code_incomplete" not in sheet_text and b"spec_incomplete" not in sheet_text

    first = {"P": "assercao_insuficiente", "R": "exige_violacao_da_regra"}
    fill(out / "auditoria-testes-codificador_a.xlsx", tmp_path / "a.xlsx", lambda g, i: first[g])
    fill(out / "auditoria-testes-codificador_b.xlsx", tmp_path / "b.xlsx",
         lambda g, i: first[g] if i != "T001" else ("nao_testa_regra" if g == "P" else "teste_quebrado"))
    fill(out / "auditoria-testes-codificador_a.xlsx", tmp_path / "final.xlsx", lambda g, i: first[g])
    key = json.loads((out / "chave-privada.json").read_text())
    result = audit.score(key, audit.suites_from_dir(out, key), tmp_path / "a.xlsx", tmp_path / "b.xlsx",
                         tmp_path / "final.xlsx")
    total = result["agreement"]["P"]["coded_by_both"] + result["agreement"]["R"]["coded_by_both"]
    assert total == 104
    assert sum(len(result["agreement"][g]["disagreements"]) for g in "PR") == 1
    assert result["final"]["R"]["code_incomplete"]["labels"]["exige_violacao_da_regra"] == 30
    assert result["blinding_source_guesses"] == {"unknown": 208}


@pytest.mark.parametrize("source_group,target_group", [("P", "R"), ("R", "P")])
def test_score_rejects_adjudicated_item_moved_to_another_group(tmp_path, source_group, target_group):
    from openpyxl import load_workbook

    frame = json.loads((audit.AUDIT / "audit-frame.json").read_text())
    out = tmp_path / "audit"
    audit.build_sheet(fake_study(tmp_path, frame), out, frame)
    labels = {"P": "assercao_insuficiente", "R": "exige_violacao_da_regra"}
    template = out / "auditoria-testes-codificador_a.xlsx"
    fill(template, tmp_path / "a.xlsx", lambda g, i: labels[g])
    fill(template, tmp_path / "final.xlsx", lambda g, i: labels[g])
    wb = load_workbook(tmp_path / "final.xlsx")
    source, target = wb[audit.SHEETS[source_group]], wb[audit.SHEETS[target_group]]
    header = [cell.value for cell in source[1]]
    values = [cell.value for cell in source[2]]
    values[header.index("categoria")] = labels[target_group]
    source.delete_rows(2)
    row_by_header = dict(zip(header, values))
    target.append([row_by_header.get(cell.value) for cell in target[1]])
    wb.save(tmp_path / "final.xlsx")
    key = json.loads((out / "chave-privada.json").read_text())
    with pytest.raises(ValueError, match="adjudicated.*wrong group"):
        audit.score(key, audit.suites_from_dir(out, key), tmp_path / "a.xlsx",
                    tmp_path / "a.xlsx", tmp_path / "final.xlsx")


def test_score_requires_literal_quote(tmp_path):
    pytest.importorskip("openpyxl")
    frame = json.loads((audit.AUDIT / "audit-frame.json").read_text())
    study = fake_study(tmp_path, frame)
    out = tmp_path / "audit"
    audit.build_sheet(study, out, frame)
    fill(out / "auditoria-testes-codificador_a.xlsx", tmp_path / "a.xlsx",
         lambda g, i: "espera_violacao" if g == "P" else "teste_quebrado", quote="not in the suite")
    key = json.loads((out / "chave-privada.json").read_text())
    with pytest.raises(ValueError, match="literal quote"):
        audit.score(key, audit.suites_from_dir(out, key), tmp_path / "a.xlsx", tmp_path / "a.xlsx", None)


def test_sheet_refuses_private_results_that_differ(tmp_path):
    pytest.importorskip("openpyxl")
    frame = json.loads((audit.AUDIT / "audit-frame.json").read_text())
    study = fake_study(tmp_path, frame)
    data = json.loads((study / "results.json").read_text())
    data["rows"][0]["verdict"] = "quiet" if data["rows"][0]["verdict"] != "quiet" else "error_alarm"
    (study / "results.json").write_text(json.dumps(data))
    with pytest.raises(ValueError, match="differ"):
        audit.build_sheet(study, tmp_path / "audit", frame)


def test_calibration_frame_is_seeded_and_balanced():
    results, manifest = published()
    first = audit.calibration_frame(results, manifest, "gpt-6-astra")
    assert first == audit.calibration_frame(results, manifest, "gpt-6-astra")
    assert [i["group"] for i in first["items"]] == ["P"] * 3 + ["R"] * 3
    assert first["scored"] is False


def test_committed_calibration_frame_matches_opus_run():
    path = audit.ROOT / "data/shared-omission-e2e/claude-evaluation-v1/claude-opus-4-6"
    if not (path / "results.json").exists():
        pytest.skip("Opus 4.6 results not on this branch yet (PR #190)")
    committed = json.loads((audit.AUDIT / "calibration-frame.json").read_text())
    rebuilt = audit.calibration_frame(json.loads((path / "results.json").read_text()),
                                      json.loads((path / "frozen-manifest-public.json").read_text()),
                                      "claude-opus-4-6")
    assert rebuilt == committed


def test_calibration_sheet_uses_its_own_ids(tmp_path):
    pytest.importorskip("openpyxl")
    results, manifest = published()
    cal = audit.calibration_frame(results, manifest, "gpt-6-astra")
    study = fake_study(tmp_path, {"items": cal["items"]})
    built = audit.build_sheet(study, tmp_path / "cal", cal, prefix="C", stem="calibracao")
    assert built["items"] == 6
    assert (tmp_path / "cal" / "calibracao-codificador_a.xlsx").exists()
    key = json.loads((tmp_path / "cal" / "chave-privada.json").read_text())
    assert all(i["item_id"].startswith("C") for i in key["items"])
