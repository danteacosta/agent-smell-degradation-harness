"""Blind human audit of a random 20% of the confirmatory screening decisions.

Pre-registration section 3 (reviewers): "a human audit of a random 20% of
panel decisions is planned before any confirmatory claim". This covers the two
confirmatory screening panels of 2026-10-05 (reserve and 2024 window).

  build  draw the sample and write a blind .xlsx: the auditor sees exactly
         what the panel saw (project, file, removed and added sentences) and
         the same three criteria; never the panel's decision, reason, target
         rule or route. Only the drawn candidate ids and the seed are written
         to the manifest, so the stratum (which encodes the decision) is not
         revealed; score recomputes it.
  design publish the stratum table (sizes, draws, inclusion probabilities and
         weights, no candidate ids) and check it reproduces the frozen sample
  score  compare the filled sheet with the panel's consensus: unweighted
         agreement and Cohen's kappa by option; agreement only (no kappa, the
         panel label is constant) within each panel decision; and a stratified
         weighted estimate of agreement over all decided candidates.

Sampling: the 365 decided candidates (2 unresolved are not decisions and
are left out) are stratified by option x project x panel decision; the
ceil(20%) = 73 draws are allocated proportionally by largest remainder, at
least one per non-empty stratum when the total allows, and drawn with a
fixed seed inside each stratum.
"""
from __future__ import annotations

import argparse
from collections import Counter
import json
import math
from pathlib import Path
import random
import statistics
import sys

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

OPTIONS = {
    "reserve": (ROOT / "data/llm-screening-confirmatory-reserve-20261005/results.json",
                ROOT / "data/requirement-sampling/screening-confirmatory-reserve-20261005.json"),
    "w2024": (ROOT / "data/llm-screening-confirmatory-w2024-20261005/results.json",
              ROOT / "data/requirement-sampling/screening-confirmatory-w2024-20261005.json"),
}
SEED = 2026100505
FRACTION = 0.20
OUT_DIR = ROOT / "data/human-audit"
SHEET = OUT_DIR / "screening-audit-sheet-20261005.xlsx"
MANIFEST = OUT_DIR / "screening-audit-sample-20261005.json"
DESIGN = OUT_DIR / "screening-audit-design-20261005.json"
REASONS = ("admit", "operations_or_configuration", "not_web_interface", "no_rule_change",
           "no_single_testable_rule", "needs_multiple_systems")

INSTRUCTIONS = [
    "Auditoria humana cega de 20% da triagem confirmatória (05/10)",
    "",
    "Não abra data/llm-screening-confirmatory-*/ antes de terminar: as decisões do painel não estão nesta planilha.",
    "Cada linha é uma mudança na documentação de usuário de um projeto: frases removidas (-) e acrescentadas (+).",
    "Decida se a mudança contém uma regra de comportamento que atende aos TRÊS critérios:",
    "1. Um usuário consegue observá-la na interface WEB do produto. Não contam instalação, configuração de",
    "   servidor ou autenticação, linha de comando, só API, recursos só mobile ou desktop, nem cenários com vários servidores.",
    "2. Pode ser escrita como UMA obrigação testável, com uma observação que passa e uma que falha.",
    "3. Uma página web pequena e autocontida, imitando aquela tela, conseguiria exercitá-la sem o produto inteiro.",
    "",
    "decisao: admit ou exclude. motivo: 'admit' exatamente quando decisao = admit; senão o critério que falhou:",
    "  operations_or_configuration, not_web_interface, no_rule_change, no_single_testable_rule, needs_multiple_systems.",
    "regra_alvo: se admit, a regra em uma frase. Comentário: dúvidas ou casos de fronteira.",
]


def load_options() -> list[dict]:
    items = []
    for option, (results, sample) in OPTIONS.items():
        texts = {c["candidate_id"]: c for c in json.loads(sample.read_text())["candidates"]}
        for r in json.loads(results.read_text())["rows"]:
            if r["consensus"] not in ("admit", "exclude"):
                continue
            c = texts[r["candidate_id"]]
            items.append({"option": option, "candidate_id": r["candidate_id"], "project": c["project"],
                          "file": c["file"], "removed": c["removed"], "added": c["added"],
                          "panel": r["consensus"], "panel_rule": r.get("target_rule")})
    return items


def allocate(sizes: dict, total: int) -> dict:
    n = sum(sizes.values())
    quota = {k: total * v / n for k, v in sizes.items()}
    alloc = {k: math.floor(q) for k, q in quota.items()}
    for k in sorted(quota, key=lambda k: (-(quota[k] - alloc[k]), k))[:total - sum(alloc.values())]:
        alloc[k] += 1
    # at least one per stratum when possible, taken from the largest allocations
    for k in sorted(sizes):
        if alloc[k] == 0 and sizes[k] > 0:
            donor = max((d for d in alloc if alloc[d] > 1), key=lambda d: (alloc[d], d), default=None)
            if donor is None:
                break
            alloc[donor] -= 1
            alloc[k] = 1
    return alloc


def strata_of(items: list[dict]) -> dict[tuple, list[dict]]:
    strata: dict[tuple, list[dict]] = {}
    for it in items:
        strata.setdefault((it["option"], it["project"], it["panel"]), []).append(it)
    return strata


def design_table(items: list[dict]) -> list[dict]:
    """Stratum sizes, draws and inclusion probabilities (no candidate ids)."""
    strata = strata_of(items)
    alloc = allocate({k: len(v) for k, v in strata.items()}, math.ceil(FRACTION * len(items)))
    return [{"option": o, "project": p, "panel_decision": d, "population": len(strata[(o, p, d)]),
             "drawn": alloc[(o, p, d)], "inclusion_probability": round(alloc[(o, p, d)] / len(strata[(o, p, d)]), 6),
             "weight": round(len(strata[(o, p, d)]) / alloc[(o, p, d)], 6)}
            for (o, p, d) in sorted(strata)]


def draw(items: list[dict]) -> list[dict]:
    strata = strata_of(items)
    total = math.ceil(FRACTION * len(items))
    alloc = allocate({k: len(v) for k, v in strata.items()}, total)
    rng = random.Random(SEED)
    chosen = []
    for key in sorted(strata):
        pool = sorted(strata[key], key=lambda i: i["candidate_id"])
        chosen.extend(rng.sample(pool, alloc[key]))
    return chosen


def build() -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    if SHEET.exists() or MANIFEST.exists():
        raise FileExistsError("the screening audit sample is drawn once")
    items = load_options()
    chosen = draw(items)
    order = chosen[:]
    random.Random(f"{SEED}-order").shuffle(order)
    wb = Workbook()
    ws = wb.active
    ws.title = "instrucoes"
    for i, line in enumerate(INSTRUCTIONS, start=1):
        ws.cell(i, 1, line).font = Font(bold=i == 1)
    ws.column_dimensions["A"].width = 140
    s = wb.create_sheet("triagem")
    s.append(["ordem", "id", "projeto", "arquivo", "removidas (-)", "acrescentadas (+)",
              "decisao", "motivo", "regra_alvo", "comentario"])
    for c in s[1]:
        c.font = Font(bold=True)
    for n, it in enumerate(order, start=1):
        s.append([n, it["candidate_id"], it["project"], it["file"], "\n".join(f"- {x}" for x in it["removed"]),
                  "\n".join(f"+ {x}" for x in it["added"]), None, None, None, None])
    for col, choices in (("G", ("admit", "exclude")), ("H", REASONS)):
        dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        s.add_data_validation(dv)
        dv.add(f"{col}2:{col}{len(order) + 1}")
    yellow, wrap = PatternFill("solid", fgColor="FFF2CC"), Alignment(wrap_text=True, vertical="top")
    for row in s.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
        for c in row[6:]:
            c.fill = yellow
    for col, w in zip("ABCDEFGHIJ", (7, 18, 13, 40, 60, 60, 12, 26, 50, 40)):
        s.column_dimensions[col].width = w
    s.freeze_panes = "C2"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(SHEET)
    manifest = {"schema_version": "screening-human-audit-sample/v1", "seed": SEED, "fraction": FRACTION,
                "population": len(items), "sample": len(chosen),
                "stratification": "option x project x panel decision; largest remainder; >=1 per stratum",
                "candidate_ids": sorted(it["candidate_id"] for it in chosen),
                "sources": {o: [str(p.relative_to(ROOT)) for p in paths] for o, paths in OPTIONS.items()}}
    MANIFEST.write_text(json.dumps(manifest, indent=2) + "\n")
    return {"sheet": str(SHEET.relative_to(ROOT)), "population": len(items), "sample": len(chosen),
            "by_option": dict(Counter(it["option"] for it in chosen))}


def _agreement(pairs):
    return round(sum(h == p for h, p in pairs) / len(pairs), 3) if pairs else None


def weighted_agreement(answered: list[tuple], panel: dict, table: list[dict]) -> dict:
    """Stratified (ratio) estimate of human-panel agreement over all decided candidates.

    Each answered item carries weight N_h / n_h of its stratum. With an incomplete
    sheet, weights use the answered count of the stratum; strata with no answer
    are reported as missing, and the estimate covers only the answered strata.
    """
    by_stratum: dict[tuple, list[int]] = {}
    for cid, human, _, _ in answered:
        it = panel[cid]
        by_stratum.setdefault((it["option"], it["project"], it["panel"]), []).append(int(human == it["panel"]))
    sizes = {(r["option"], r["project"], r["panel_decision"]): r["population"] for r in table}
    out = {}
    for scope in ("reserve", "w2024", "all"):
        keys = [k for k in sizes if scope == "all" or k[0] == scope]
        covered = [k for k in keys if by_stratum.get(k)]
        total = sum(sizes[k] for k in covered)
        if not total:
            out[scope] = {"estimate": None}
            continue
        est = sum(sizes[k] * (sum(by_stratum[k]) / len(by_stratum[k])) for k in covered) / total
        # stratified variance with finite-population correction; strata with one answer add none
        var = sum((sizes[k] / total) ** 2 * (1 - len(by_stratum[k]) / sizes[k])
                  * (statistics.variance(by_stratum[k]) / len(by_stratum[k]))
                  for k in covered if len(by_stratum[k]) > 1)
        out[scope] = {"estimate": round(est, 4), "se_lower_bound": round(math.sqrt(var), 4),
                      "population_covered": total, "population": sum(sizes[k] for k in keys),
                      "strata_missing": len(keys) - len(covered),
                      "strata_single_answer": sum(1 for k in covered if len(by_stratum[k]) == 1)}
    out["note"] = ("se_lower_bound omits strata with a single answer (no within-stratum variance), so the true "
                   "standard error is larger.")
    return out


def score(sheet_path: Path) -> dict:
    from openpyxl import load_workbook
    from scripts.llm_screening_panel import cohen_kappa
    panel = {it["candidate_id"]: it for it in load_options()}
    rows = list(load_workbook(sheet_path, read_only=True)["triagem"].iter_rows(values_only=True))
    header = list(rows[0])
    frozen_ids = set(json.loads(MANIFEST.read_text())["candidate_ids"])
    seen_ids = set()
    answered, bad = [], []
    for r in rows[1:]:
        rec = dict(zip(header, r))
        if not rec.get("id"):
            continue
        candidate_id = rec["id"]
        if candidate_id not in frozen_ids:
            raise ValueError(f"candidate {candidate_id!r} outside frozen audit sample")
        if candidate_id in seen_ids:
            raise ValueError(f"duplicate audit candidate: {candidate_id!r}")
        seen_ids.add(candidate_id)
        decision = str(rec.get("decisao") or "").strip().lower()
        reason = str(rec.get("motivo") or "").strip()
        if not decision:
            continue
        if decision not in ("admit", "exclude") or reason not in REASONS or (decision == "admit") != (reason == "admit"):
            bad.append(rec["id"])
            continue
        answered.append((rec["id"], decision, reason, rec.get("regra_alvo")))

    def pairs_for(option=None, decision=None):
        return [(h, panel[i]["panel"]) for i, h, _, _ in answered
                if (option is None or panel[i]["option"] == option) and (decision is None or panel[i]["panel"] == decision)]

    def kappa(pairs):
        return cohen_kappa([h for h, _ in pairs], [p for _, p in pairs]) if pairs else None

    table = json.loads(DESIGN.read_text())["strata"]
    return {
        "answered": len(answered), "of": len(json.loads(MANIFEST.read_text())["candidate_ids"]),
        "malformed": bad,
        "sample_unweighted": {
            "agreement": _agreement(pairs_for()),
            "kappa_by_option": {o: kappa(pairs_for(option=o)) for o in ("reserve", "w2024")},
            "kappa_pooled_descriptive": kappa(pairs_for()),
            "note": "unweighted over the stratified sample; strata are not sampled in proportion to their size",
        },
        "agreement_within_panel_decision": {
            d: {"n": len(pairs_for(decision=d)), "agreement": _agreement(pairs_for(decision=d)),
                "note": "the panel label is constant here, so only agreement is reported, not kappa"}
            for d in ("admit", "exclude")},
        "weighted_agreement_full_population": weighted_agreement(answered, panel, table),
        "disagreements": [{"candidate_id": i, "project": panel[i]["project"], "option": panel[i]["option"],
                           "human": h, "human_reason": reason, "panel": panel[i]["panel"]}
                          for i, h, reason, _ in answered if h != panel[i]["panel"]],
        "both_admit_rules": [{"candidate_id": i, "human_rule": rule, "panel_rule": panel[i]["panel_rule"]}
                             for i, h, _, rule in answered if h == "admit" == panel[i]["panel"]],
    }


def write_design() -> dict:
    """Publish the stratum table and check it reproduces the frozen sample."""
    items = load_options()
    frozen = sorted(json.loads(MANIFEST.read_text())["candidate_ids"])
    if sorted(it["candidate_id"] for it in draw(items)) != frozen:
        raise ValueError("the sampling procedure no longer reproduces the frozen sample")
    table = design_table(items)
    doc = {"schema_version": "screening-human-audit-design/v1", "seed": SEED, "fraction": FRACTION,
           "population": len(items), "sample": sum(r["drawn"] for r in table),
           "estimator": "stratified ratio estimate: sum_h N_h * mean_h / sum_h N_h over answered strata",
           "strata": table}
    DESIGN.write_text(json.dumps(doc, indent=2) + "\n")
    return {"strata": len(table), "min_inclusion_probability": min(r["inclusion_probability"] for r in table),
            "max_inclusion_probability": max(r["inclusion_probability"] for r in table)}


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="mode", required=True)
    sub.add_parser("build")
    sub.add_parser("score").add_argument("sheet", type=Path)
    sub.add_parser("design")
    args = parser.parse_args(argv)
    result = {"build": build, "design": write_design}.get(args.mode, lambda: score(args.sheet))()
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
