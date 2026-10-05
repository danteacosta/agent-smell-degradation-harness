"""Blind human audit of the LLM-panel decisions (pre-registration section 8, decision 2).

  build  write an .xlsx with the exact inputs each panel saw and empty answer
         columns. The panel's decisions are NOT in the sheet; they stay in the
         published results files, which the auditor should not open first.
  score  read the filled sheet and compare it with the panel's final decisions:
         agreement, Cohen's kappa and the list of disagreements.

Two audits:
  context_cue  46 cases (data/context-cue/20261004-v2/results.json)
  historical   40 old-documentation reviews (data/historical-arm/panel/results.json)
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import random
import sys

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

CUE_RESULTS = ROOT / "data/context-cue/20261004-v2/results.json"
HIST_RESULTS = ROOT / "data/historical-arm/panel/results.json"
HIST_CLASSIFICATION = ROOT / "data/historical-arm/classification.json"
SEED = 2026100504

INSTRUCTIONS = [
    "Auditoria humana cega das decisões dos painéis de LLM",
    "",
    "Não abra os arquivos de resultados dos painéis antes de terminar: as decisões dos modelos não estão nesta planilha de propósito.",
    "Preencha só as colunas amarelas. Use as listas suspensas. Citação: copie o trecho exato do texto que justifica a resposta.",
    "Se não tiver certeza, responda mesmo assim e explique na coluna de comentário.",
    "",
    "ABA context_cue (46 linhas)",
    "Um modelo recebeu o TEXTO e teve de implementar uma página. A REGRA OMITIDA foi retirada do pedido antes.",
    "Pergunta: a regra omitida, ou uma consequência observável dela, pode ser inferida só a partir do TEXTO?",
    "Conte outras frases do pedido, rótulos, ids, nomes de campo, dados e nomes de API que estão no TEXTO.",
    "Não conte conhecimento geral sobre o produto.",
    "Responda 'sim' só se o TEXTO afirma ou implica a condição E o comportamento exigido pela regra.",
    "Não basta conter os elementos de que a regra fala (uma lista de subtarefas não diz o que acontece com elas).",
    "Se o trecho citado continuaria verdadeiro numa página que viola a regra, responda 'não'.",
    "",
    "ABA historico (40 linhas)",
    "FEATURE: o pedido sem a regra. RULE: a regra na versão nova. OLD DOCUMENTATION: o trecho da documentação antes do commit.",
    "1. funcionalidade_documentada: a documentação antiga já descreve a funcionalidade do FEATURE (sim/não)?",
    "   Uma funcionalidade diferente na mesma página não conta.",
    "2. status_da_regra na documentação antiga:",
    "   same: afirma a regra com o mesmo sentido;",
    "   vaguer: afirma o comportamento de forma menos precisa, compatível com a regra mas sem determiná-la;",
    "   absent: não diz nada sobre esse comportamento;",
    "   different: afirma um comportamento que contradiz a regra (o produto mudou).",
    "Use só a documentação antiga mostrada, sem conhecimento do produto.",
]


def cue_items() -> list[dict]:
    from scripts.context_cue_panel import selected_cases
    return [{"id": i["case"], "regra_omitida": i["omitted"], "texto": i["text"]} for i in selected_cases()]


def hist_items() -> list[dict]:
    rows = json.loads(HIST_CLASSIFICATION.read_text())["rows"]
    return [{"id": r["case"], "feature": r["requirement_without_rule"], "rule": r["c_span"].strip(),
             "old_documentation": r["old_excerpt"]} for r in rows if r["class"] == "panel"]


def build(out: Path) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    if out.exists():
        raise FileExistsError(f"{out} exists; the audit sheet is built once")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")
    wb = Workbook()
    ws = wb.active
    ws.title = "instrucoes"
    for i, line in enumerate(INSTRUCTIONS, start=1):
        ws.cell(i, 1, line).font = Font(bold=i == 1)
    ws.column_dimensions["A"].width = 140

    def sheet(title, items, text_cols, answer_cols):
        s = wb.create_sheet(title)
        order = items[:]
        random.Random(f"{SEED}-{title}").shuffle(order)
        header = ["ordem", "id", *text_cols, *[a for a, _ in answer_cols], "comentario"]
        s.append(header)
        for c in s[1]:
            c.font = Font(bold=True)
        for n, item in enumerate(order, start=1):
            s.append([n, item["id"], *[item[c] for c in text_cols], *[None] * (len(answer_cols) + 1)])
        first_answer = 3 + len(text_cols)
        for j, (name, choices) in enumerate(answer_cols):
            col = s.cell(1, first_answer + j).column_letter
            if choices:
                dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
                s.add_data_validation(dv)
                dv.add(f"{col}2:{col}{len(order) + 1}")
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
            for c in row[first_answer - 1:]:
                c.fill = yellow
        widths = [7, 34, *[70] * len(text_cols), *[18] * len(answer_cols), 40]
        for i, w in enumerate(widths, start=1):
            s.column_dimensions[s.cell(1, i).column_letter].width = w
        s.freeze_panes = "C2"
        return len(order)

    n_cue = sheet("context_cue", cue_items(), ["regra_omitida", "texto"],
                  [("context_cue", ["sim", "não"]), ("citacao", None)])
    n_hist = sheet("historico", hist_items(), ["feature", "rule", "old_documentation"],
                   [("funcionalidade_documentada", ["sim", "não"]),
                    ("status_da_regra", ["same", "vaguer", "absent", "different"]),
                    ("citacao_funcionalidade", None), ("citacao_regra", None)])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"file": str(out), "context_cue": n_cue, "historico": n_hist}


def kappa(a: list[str], b: list[str]) -> float | None:
    from scripts.llm_screening_panel import cohen_kappa
    return cohen_kappa(a, b) if a else None


def score(sheet_path: Path) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(sheet_path, read_only=True)

    def answers(title):
        rows = list(wb[title].iter_rows(values_only=True))
        header = list(rows[0])
        return {r[header.index("id")]: dict(zip(header, r)) for r in rows[1:] if r[header.index("id")]}

    yes = {"sim": "yes", "não": "no", "nao": "no"}
    cue_panel = {r["case"]: ("yes" if r["context_cue"] == 1 else "no" if r["context_cue"] == 0 else None)
                 for r in json.loads(CUE_RESULTS.read_text())["rows"]}
    human = {k: yes.get(str(v.get("context_cue") or "").strip().lower()) for k, v in answers("context_cue").items()}
    pairs = [(human[k], cue_panel[k]) for k in sorted(human) if human[k] and cue_panel.get(k)]
    out = {"context_cue": {
        "answered": len(pairs), "of": len(cue_panel),
        "agreement": round(sum(h == p for h, p in pairs) / len(pairs), 3) if pairs else None,
        "kappa": kappa([h for h, _ in pairs], [p for _, p in pairs]),
        "disagreements": [{"case": k, "human": human[k], "panel": cue_panel[k]}
                          for k in sorted(human) if human[k] and cue_panel.get(k) and human[k] != cue_panel[k]]}}

    final = {r["case"]: r["final"] for r in json.loads(HIST_RESULTS.read_text())["rows"]}
    hist = answers("historico")
    rows = []
    for k, v in sorted(hist.items()):
        f = yes.get(str(v.get("funcionalidade_documentada") or "").strip().lower())
        s = str(v.get("status_da_regra") or "").strip().lower() or None
        p = final.get(k)
        if f and s and p:
            rows.append((k, f, s, p["feature_documented"], p["rule_status"]))

    def admitted(f, s):
        return f == "yes" and s in ("vaguer", "absent")

    out["historical"] = {
        "answered": len(rows), "of": len(final),
        "kappa_feature_documented": kappa([r[1] for r in rows], [r[3] for r in rows]),
        "kappa_rule_status": kappa([r[2] for r in rows], [r[4] for r in rows]),
        "admission_agreement": round(sum(admitted(r[1], r[2]) == admitted(r[3], r[4]) for r in rows) / len(rows), 3)
        if rows else None,
        "disagreements": [{"case": r[0], "human": [r[1], r[2]], "panel": [r[3], r[4]],
                           "admission_changes": admitted(r[1], r[2]) != admitted(r[3], r[4])}
                          for r in rows if (r[1], r[2]) != (r[3], r[4])]}
    return out


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="mode", required=True)
    b = sub.add_parser("build")
    b.add_argument("--out", type=Path, default=ROOT / "data/human-audit/audit-sheet-20261005.xlsx")
    s = sub.add_parser("score")
    s.add_argument("sheet", type=Path)
    args = parser.parse_args(argv)
    result = build(args.out) if args.mode == "build" else score(args.sheet)
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
